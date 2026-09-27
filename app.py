import streamlit as st
import google.generativeai as genai
import fitz  # PyMuPDF
from docx import Document
from docx.shared import RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import time
import os
import io
import re
from datetime import datetime
import summary_design as sd

# Setup the page configuration
st.set_page_config(page_title="PDF Read Tool", page_icon="📄", layout="centered")

st.title("📄 PDF Read Tool - Smart Analyzer")
st.write("Apni PDF yahan upload karein aur AI uski summary nikal dega.")

# Sidebar for configuration
with st.sidebar:
    st.header("Settings")
    api_key = st.text_input("Google Gemini API Key", type="password")
    start_page = st.number_input("Start Page", min_value=1, value=1)
    end_page = st.number_input("End Page", min_value=1, value=1000)

def add_formatted_paragraph(doc, line, style=None):
    p = doc.add_paragraph(style=style)
    chunks = re.split(r'(\*\*.*?\*\*)', line)
    for chunk in chunks:
        if chunk.startswith('**') and chunk.endswith('**'):
            run = p.add_run(chunk[2:-2])
            run.bold = True
        else:
            p.add_run(chunk)
    return p

def add_markdown_to_doc(doc, text):
    lines = text.split('\n')
    in_table = False
    table_headers = []
    table_rows = []
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('|') and line.endswith('|'):
            if re.match(r'^\|[\s\-\|]+\|$', line):
                continue
            row = [cell.strip().replace('**', '') for cell in line.split('|')[1:-1]]
            
            if not in_table:
                table_headers = row
                in_table = True
            else:
                table_rows.append(row)
        else:
            if in_table:
                sd.add_data_table(doc, table_headers, table_rows)
                in_table = False
                table_headers = []
                table_rows = []
                
            if line.startswith('# '):
                sd.add_section_heading(doc, line[2:].replace('**', '').strip())
            elif line.startswith('## '):
                sd.add_section_heading(doc, line[3:].replace('**', '').strip())
            elif line.startswith('### '):
                sd.add_section_heading(doc, line[4:].replace('**', '').strip())
            elif line.startswith('- ') or line.startswith('* '):
                clean_line = line[2:].strip()
                add_formatted_paragraph(doc, clean_line, style='List Bullet')
            elif line == '---' or line == '***':
                continue # Skip raw markdown dividers
            else:
                add_formatted_paragraph(doc, line)
                
    if in_table:
        sd.add_data_table(doc, table_headers, table_rows)

def analyze_page(model, text, page_num):
    prompt = f"""You are an Expert Financial Analyst specializing in the Pakistan Stock Exchange (PSX).

Your objective is to read this single page of a financial report and extract ONLY new, critical information. 
DO NOT output boilerplate company introductions if they are already obvious. DO NOT invent information.
If the page contains only empty space, signatures, or irrelevant filler text, reply with exactly "EMPTY_PAGE".

WHAT TO EXTRACT (If present on this page):
1. New Developments: Any new projects, partnerships, or strategic changes.
2. Financial Numbers: Extract revenue, profit, margins, EPS, etc. YOU MUST FORMAT THESE NUMBERS AS A MARKDOWN TABLE. Add a column for YoY Growth % if possible. Use + for positive growth and - for negative growth.
3. Management Commentary: Real outlook, future plans, risks, or challenges.

STRICT RULES:
1. Language: Roman Urdu (easy words). Use English for financial terms (Revenue, EPS).
2. No Repetition: Do not say "Here is the summary". Just give the facts.
3. Tables: Whenever you see financial figures (e.g. 2024 vs 2025), put them in a Markdown Table format like this:
| Indicator | 2024 | 2025 | Growth % |
|---|---|---|---|
| Revenue | 100 | 120 | +20% |
4. Headings: Use ## for topics. Do not use # (H1).
5. Bold Text: Use **bold** for key metrics, names, and important highlights.

PAGE TEXT:
{text}"""
    try:
        response = model.generate_content(prompt)
        if "EMPTY_PAGE" in response.text:
            return ""
        return response.text
    except Exception as e:
        return f"Error on page {page_num}: {str(e)}"

uploaded_file = st.file_uploader("PDF File upload karein", type=["pdf"])

if st.button("Start Analysis"):
    if not uploaded_file:
        st.error("Bhai pehlay PDF file toh upload karein!")
    elif not api_key:
        st.error("API Key zaroori hai!")
    else:
        try:
            genai.configure(api_key=api_key)
            model = genai.GenerativeModel("gemini-flash-lite-latest")
            
            pdf_bytes = uploaded_file.read()
            pdf_document = fitz.open(stream=pdf_bytes, filetype="pdf")
            
            total_pages = len(pdf_document)
            end_page_to_process = min(end_page, total_pages)
            
            progress_bar = st.progress(0)
            status_text = st.empty()
            
            total_pages_to_process = end_page_to_process - (start_page - 1)
            pages_done = 0
            full_extracted_text = ""
            
            for page_num in range(start_page - 1, end_page_to_process):
                status_text.text(f"Page {page_num + 1} analyze ho rahi hai... Please wait.")
                
                page = pdf_document.load_page(page_num)
                text = page.get_text("text")
                
                if text.strip():
                    summary = analyze_page(model, text, page_num + 1)
                    if summary and summary.strip():
                        full_extracted_text += f"\n--- Page {page_num + 1} ---\n{summary}\n"
                
                pages_done += 1
                progress_bar.progress(pages_done / total_pages_to_process)
                
                if page_num < end_page_to_process - 1:
                    time.sleep(5)
            
            final_summary = ""
            company_name = "Financial Analysis Report"
            report_title = "AI Executive Summary"
            kpi_cards = []
            
            if full_extracted_text.strip():
                status_text.text("Generating Final Executive Summary & KPIs...")
                final_prompt = f"""You are an Expert Financial Analyst. Based on the following extracted data, write a comprehensive 'Executive Summary' in Roman Urdu.

REQUIREMENTS:
1. Identify the Company Name and a short Report Title (e.g. "Interim Results 2026").
2. Extract the 3 to 4 most critical financial metrics (Revenue, Profit, etc.) for KPI Cards. Formatted exactly as a JSON array.
3. Write the Executive Summary in Roman Urdu using short bullet points (- ) for highlights. DO NOT write thick paragraphs. Keep it punchy and professional.

STRICT OUTPUT FORMAT MUST BE EXACTLY LIKE THIS:
COMPANY_NAME: [Company Name]
REPORT_TITLE: [Report Title]

```json
[
  {{"title": "Revenue", "value": "Rs. 49,716M", "change": "35.3%", "positive": true}},
  {{"title": "Net Profit", "value": "Rs. 6,050M", "change": "17.4%", "positive": true}}
]
```

## Overall Performance
(your text here)

## Key Highlights
(your text here)

## Future Outlook
(your text here)

EXTRACTED DATA:
{full_extracted_text}"""
                final_response = model.generate_content(final_prompt)
                if final_response.text:
                    resp_text = final_response.text
                    
                    # Parse Company Name
                    cn_match = re.search(r'COMPANY_NAME:\s*(.+)', resp_text)
                    if cn_match: company_name = cn_match.group(1).strip()
                        
                    # Parse Report Title
                    rt_match = re.search(r'REPORT_TITLE:\s*(.+)', resp_text)
                    if rt_match: report_title = rt_match.group(1).strip()
                    
                    # Parse JSON KPIs
                    json_match = re.search(r'```json\s*(.*?)\s*```', resp_text, re.DOTALL)
                    if json_match:
                        import json
                        try:
                            kpi_cards = json.loads(json_match.group(1))
                        except:
                            pass
                    
                    # Clean Markdown text
                    clean_md = re.sub(r'COMPANY_NAME:.*?\n', '', resp_text)
                    clean_md = re.sub(r'REPORT_TITLE:.*?\n', '', clean_md)
                    clean_md = re.sub(r'```json.*?```', '', clean_md, flags=re.DOTALL)
                    final_summary = clean_md.strip()

            status_text.text("Analysis Complete! Generating File...")
            
            doc = Document()
            sd.add_footer_page_number(doc.sections[0])
            
            # --- PROFESSIONAL COVER PAGE ---
            doc.add_heading(company_name.upper(), 0)
            p_sub = doc.add_paragraph(report_title)
            p_sub.runs[0].font.size = Pt(14)
            p_sub.runs[0].font.color.rgb = sd.NAVY
            p_sub.runs[0].bold = True
            
            doc.add_paragraph(f"Generated on: {datetime.now().strftime('%d %B %Y')}  ·  Source: {uploaded_file.name}")
            
            # Add KPI Cards right below header if available
            if kpi_cards:
                doc.add_paragraph() # spacer
                sd.add_kpi_row(doc, kpi_cards)
                doc.add_paragraph() # spacer
            
            # Add AI Summary Markdown
            if final_summary:
                add_markdown_to_doc(doc, final_summary)
                doc.add_page_break()
            
            # --- APPENDIX: RAW EXTRACTS ---
            sd.add_section_heading(doc, "Appendix: Detailed Page Extracts")
            for line in full_extracted_text.split('\n'):
                if line.startswith('--- Page ') and line.endswith(' ---'):
                    match = re.search(r'Page (\d+)', line)
                    if match:
                        p = doc.add_paragraph(f"Source: Page {match.group(1)}")
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        p.runs[0].font.size = Pt(9)
                        p.runs[0].font.color.rgb = RGBColor(128, 128, 128)
                else:
                    if line.strip():
                        add_markdown_to_doc(doc, line)
            
            doc_io = io.BytesIO()
            doc.save(doc_io)
            doc_io.seek(0)
            
            st.success("✅ Zabardast! File tayar hai.")
            
            file_name = f"{os.path.splitext(uploaded_file.name)[0]} - Summary.docx"
            st.download_button(
                label="📥 Download Professional Report (.docx)",
                data=doc_io,
                file_name=file_name,
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            
        except Exception as e:
            st.error(f"Koi masla aagaya: {str(e)}")
