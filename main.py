import google.generativeai as genai
import fitz  # PyMuPDF
from docx import Document
import re
import time
import os
import sys

try:
    with open(".api_key", "r") as f:
        API_KEY = f.read().strip()
except:
    API_KEY = ""
START_PAGE = 1
END_PAGE = 1000

genai.configure(api_key=API_KEY)
model = genai.GenerativeModel("gemini-flash-lite-latest")

def add_markdown_to_doc(doc, text):
    lines = text.split('\n')
    for line in lines:
        if line.startswith('# '):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(line[2:].strip(), style='List Bullet')
        elif line.strip() != "":
            # Basic text paragraph (skipping complex tables for simplicity, treating as text)
            # You can enhance table parsing if needed, but for now paragraphs are fine.
            clean_line = line.replace('**', '') # strip bold stars for neatness
            doc.add_paragraph(clean_line.strip())

def analyze_page(text, page_num):
    prompt = f"""You are an Expert Investor, a First-Principles Business Operator, and a Financial Analyst specializing in the Pakistan Stock Exchange (PSX). You have built, scaled, and exited successful companies, so you know how to look past the marketing fluff and identify the real core of a business.

Your objective is to read the provided text from a company's financial statement / report carefully and extract the absolute most critical information that an investor MUST know.

WHAT TO EXTRACT:
1. Company Introduction: What the company does, its core business, sector, and any background context provided.
2. New Developments: Any new advancements, expansions, projects, partnerships, product launches, or strategic changes mentioned by the company.
3. Financial Results: Revenue, profit/loss, margins, EPS, growth or decline compared to previous period, and any other key financial figures reported.
4. Management Commentary: Everything the company wants to communicate to its investors — outlook, future plans, challenges, risks, and management's own remarks or explanations.
5. Any other material fact an investor should know before making a decision.

STRICT RULES FOR YOUR RESPONSE:
1. Language: Your entire output MUST be in Roman Urdu. Use simple, easy-to-understand words (asan alfaz) so a common retail investor can easily grasp complex financial concepts. Do NOT write sentences in English. You may only use English for unavoidable financial terms (like 'EBITDA', 'IPO', 'Revenue', 'EPS', etc.), but explain their context in Roman Urdu.
2. No Fluff & High Precision: Be extremely precise, authentic, and to the point. Do not add filler words or generic advice. Every single sentence must carry maximum weight.
3. Structure and Formatting: Provide neat and clean formatting. Use H1 for main section, H2 for sub-topics. Use bullet points instead of long paragraphs.
4. Focus: Extract the core business model, new developments, financial performance, risks, and real truths — everything the company is communicating to its investors.

PAGE TEXT (Page {page_num}):
{text}"""
    try:
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        print(f"Error analyzing page {page_num}: {e}")
        return ""

def main():
    print("="*50)
    
    if len(sys.argv) > 1:
        pdf_path_input = sys.argv[1].strip()
    else:
        pdf_path_input = input("PDF file ka path yahan drag & drop karein (ya naam likhein): ").strip()
        
    # Remove quotes if dragged from terminal
    pdf_path_input = pdf_path_input.strip("'").strip('"')
    
    if not pdf_path_input:
        print("Aapne koi file path nahi diya! Script band ho rahi hai.")
        return
        
    base_name = os.path.basename(pdf_path_input)
    file_name_without_ext = os.path.splitext(base_name)[0]
    output_file = f"{file_name_without_ext} - Summary.docx"

    print("\nPDF ANALYZER SHURU HO GAYA")
    
    try:
        pdf_document = fitz.open(pdf_path_input)
    except Exception as e:
        print(f"File open karne mein masla: {e}. Please ensure path is correct.")
        return

    doc = Document()
    total_pages = len(pdf_document)
    end_page_to_process = min(END_PAGE, total_pages)
    
    for page_num in range(START_PAGE - 1, end_page_to_process):
        print(f"Page {page_num + 1} analyze ho rahi hai...")
        page = pdf_document.load_page(page_num)
        text = page.get_text("text")
        
        if text.strip():
            summary = analyze_page(text, page_num + 1)
            if summary:
                doc.add_heading(f'Page {page_num + 1} Summary', level=1)
                add_markdown_to_doc(doc, summary)
                doc.add_paragraph('\n')
        
        print(f"Page {page_num + 1} done!")
        time.sleep(5) # Anti-ban rate limit protection for Free Tier
    
    doc.save(output_file)
    print(f"COMPLETE! '{output_file}' ready hai!")

if __name__ == "__main__":
    main()
